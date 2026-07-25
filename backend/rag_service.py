import os
from pathlib import Path
from typing import Any, Dict, List, Optional

import docx
import fitz  # PyMuPDF
import pandas as pd
from dotenv import load_dotenv
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_openai import ChatOpenAI
from langchain_text_splitters import RecursiveCharacterTextSplitter

from backend.config import (
    CHUNK_OVERLAP,
    CHUNK_SIZE,
    EMBEDDING_MODEL,
    FAISS_INDEX_PATH,
    LLM_MODEL,
    OPENROUTER_API_KEY,
    OPENROUTER_BASE_URL,
    TEMPERATURE,
    TOP_K,
)
from backend.vision_processor import (
    generate_image_caption,
    parse_and_caption_pdf_images,
)

# Load environment configuration
load_dotenv()

# Global vector database handler
global_vector_db: Optional[FAISS] = None


def _get_embedding_instance() -> HuggingFaceEmbeddings:
    """Helper to instantiate HuggingFace Feature Extractor for vector embeddings."""
    return HuggingFaceEmbeddings(model_name=EMBEDDING_MODEL)


def initialize_or_update_index(documents: List[Document]) -> FAISS:
    """Creates a new FAISS index or appends new documents to an existing index."""
    global global_vector_db
    embedder = _get_embedding_instance()

    if global_vector_db is None and Path(FAISS_INDEX_PATH).exists():
        try:
            global_vector_db = FAISS.load_local(
                FAISS_INDEX_PATH,
                embedder,
                allow_dangerous_deserialization=True,
            )
        except Exception:
            global_vector_db = None

    if global_vector_db is None:
        global_vector_db = FAISS.from_documents(documents, embedder)
    else:
        global_vector_db.add_documents(documents)

    global_vector_db.save_local(FAISS_INDEX_PATH)
    return global_vector_db


def retrieve_vector_store() -> FAISS:
    """Ensures index is loaded into memory and returns the active FAISS instance."""
    global global_vector_db
    if global_vector_db is None:
        embedder = _get_embedding_instance()
        global_vector_db = FAISS.load_local(
            FAISS_INDEX_PATH,
            embedder,
            allow_dangerous_deserialization=True,
        )
    return global_vector_db


def _initialize_llm() -> ChatOpenAI:
    """Constructs OpenRouter ChatOpenAI client interface."""
    return ChatOpenAI(
        model=LLM_MODEL,
        temperature=TEMPERATURE,
        base_url=OPENROUTER_BASE_URL,
        api_key=OPENROUTER_API_KEY,
    )


def format_grounded_prompt(user_query: str, retrieved_docs: List[Document]) -> str:
    """Builds a structured prompt bounding the context strictly to retrieved documents."""
    formatted_sources = []
    for doc in retrieved_docs:
        source_name = Path(doc.metadata.get("source", "Unknown")).name
        page_no = doc.metadata.get("page", 1)
        doc_type = doc.metadata.get("type", "text")
        formatted_sources.append(
            f"[Source: {source_name} | Page/Row: {page_no} | Format: {doc_type}]\n{doc.page_content}"
        )

    context_str = "\n\n---\n\n".join(formatted_sources)

    return f"""You are an intelligent multimodal document-grounded assistant.
Synthesize an accurate answer using ONLY the factual context provided below.
If the information is not present in the context, explicitly respond with:
"I don't know based on the provided document."

Context Documents:
{context_str}

User Question:
{user_query}

Answer:"""


# --- INGESTION PIPELINES ---

def ingest_pdf(file_path: str) -> Dict[str, Any]:
    """Ingests PDF text and embedded visual figures using Vision OCR."""
    pdf_doc = fitz.open(file_path)
    text_docs = [
        Document(
            page_content=page.get_text("text"),
            metadata={"source": file_path, "page": i + 1},
        )
        for i, page in enumerate(pdf_doc)
    ]
    pdf_doc.close()

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
    )
    text_chunks = splitter.split_documents(text_docs)

    # Process embedded images via vision model
    visual_records = parse_and_caption_pdf_images(file_path)
    visual_docs = [
        Document(
            page_content=item["content"],
            metadata={
                "source": file_path,
                "page": item["page"],
                "type": "embedded_image",
            },
        )
        for item in visual_records
    ]

    all_payloads = text_chunks + visual_docs
    initialize_or_update_index(all_payloads)

    return {
        "pages": len(text_docs),
        "chunks": len(text_chunks),
        "images_extracted": len(visual_docs),
        "message": "PDF document and visual figures successfully indexed.",
    }


def ingest_image(file_path: str) -> Dict[str, Any]:
    """Ingests standalone image file by converting visual information to text descriptions."""
    caption_content = generate_image_caption(file_path)

    img_doc = Document(
        page_content=caption_content,
        metadata={
            "source": file_path,
            "page": 1,
            "type": "standalone_image",
        },
    )

    initialize_or_update_index([img_doc])

    return {
        "message": "Image successfully parsed and indexed.",
        "preview": caption_content[:200],
    }


def ingest_csv(file_path: str) -> Dict[str, Any]:
    """Ingests tabular CSV data by converting structured rows into context documents."""
    dataframe = pd.read_csv(file_path)
    row_documents = []

    for row_idx, row_data in dataframe.iterrows():
        row_str = " | ".join([f"{column}: {row_data[column]}" for column in dataframe.columns])
        row_documents.append(
            Document(
                page_content=row_str,
                metadata={
                    "source": file_path,
                    "page": row_idx + 1,
                    "type": "csv_row",
                },
            )
        )

    initialize_or_update_index(row_documents)

    return {
        "total_rows": len(row_documents),
        "columns": list(dataframe.columns),
        "message": "CSV dataset successfully indexed.",
    }


def ingest_docx(file_path: str) -> Dict[str, Any]:
    """Extracts text content from Word (.docx) files and indexes chunked representations."""
    doc_file = docx.Document(file_path)
    paragraphs = [p.text for p in doc_file.paragraphs if p.text.strip()]
    full_text = "\n".join(paragraphs)

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
    )
    chunks = splitter.split_text(full_text)

    docx_docs = [
        Document(
            page_content=chunk_text,
            metadata={
                "source": file_path,
                "page": idx + 1,
                "type": "docx_chunk",
            },
        )
        for idx, chunk_text in enumerate(chunks)
    ]

    initialize_or_update_index(docx_docs)

    return {
        "chunks": len(docx_docs),
        "message": "Word document successfully indexed.",
    }


# --- QUERY EXECUTORS ---

def ask_without_rag(question: str) -> str:
    """Direct query to the LLM without context retrieval."""
    llm = _initialize_llm()
    completion = llm.invoke(question)
    return completion.content


def ask_with_rag(question: str) -> Dict[str, Any]:
    """Executes RAG flow: similarity search + context synthesis."""
    db_store = retrieve_vector_store()
    matches = db_store.similarity_search_with_score(question, k=TOP_K)

    retrieved_documents = [doc for doc, _ in matches]
    constructed_prompt = format_grounded_prompt(question, retrieved_documents)

    llm = _initialize_llm()
    generated_answer = llm.invoke(constructed_prompt)

    retrieved_chunks_meta = []
    for doc, distance_score in matches:
        retrieved_chunks_meta.append(
            {
                "content": doc.page_content,
                "source": Path(doc.metadata.get("source", "")).name,
                "page": doc.metadata.get("page", 1),
                "score": float(distance_score),
            }
        )

    return {
        "answer": generated_answer.content,
        "chunks": retrieved_chunks_meta,
    }